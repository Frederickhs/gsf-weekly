from __future__ import annotations

import threading
import tkinter as tk
from dataclasses import dataclass, field
from tkinter import messagebox, ttk

from openpyxl import load_workbook

from input_scanner import WorkbookStructure, cell_value, clean_text


DUE_SOON_DAYS = 15


@dataclass(frozen=True)
class PendingItem:
    item_title: str
    days_until_due: int | float
    status: str
    organization: str


@dataclass
class EmployeeStatus:
    employee_id: str
    employee_name: str = ""
    parks: set[str] = field(default_factory=set)
    departments: set[str] = field(default_factory=set)
    organizations: set[str] = field(default_factory=set)
    pending_items: list[PendingItem] = field(default_factory=list)

    @property
    def overall_status(self) -> str:
        if any(item.status == "Overdue" for item in self.pending_items):
            return "Overdue"
        if self.pending_items:
            return "Due Soon"
        return "Good"


def parse_days(value: object) -> int | float | None:
    if value is None or value == "":
        return None

    try:
        number = float(value)
    except (TypeError, ValueError):
        return None

    return int(number) if number.is_integer() else number


def technician_name(row: tuple, structure: WorkbookStructure) -> str:
    first_column = structure.columns.get("first_name")
    last_column = structure.columns.get("last_name")

    first_name = clean_text(cell_value(row, first_column)) if first_column else ""
    last_name = clean_text(cell_value(row, last_column)) if last_column else ""
    full_name = " ".join(part for part in (first_name, last_name) if part).strip()

    if full_name:
        return full_name

    return clean_text(cell_value(row, structure.columns["display_name"]))


def lookup_employee_status(
    structure: WorkbookStructure,
    employee_id: str,
) -> EmployeeStatus | None:
    target = employee_id.strip()
    if not target:
        return None

    workbook = load_workbook(
        filename=structure.file_path,
        read_only=True,
        data_only=True,
    )

    result: EmployeeStatus | None = None
    seen_pending: set[tuple[str, int | float, str, str]] = set()

    try:
        worksheet = workbook[structure.worksheet_name]

        for row in worksheet.iter_rows(
            min_row=structure.header_row + 1,
            values_only=True,
        ):
            row_employee_id = clean_text(
                cell_value(row, structure.columns["employee_id"])
            )
            if row_employee_id != target:
                continue

            if result is None:
                result = EmployeeStatus(employee_id=target)

            name = technician_name(row, structure)
            if name:
                result.employee_name = name

            park = clean_text(
                cell_value(row, structure.columns["park_description"])
            )
            department = clean_text(
                cell_value(row, structure.columns["department"])
            )
            organization = clean_text(
                cell_value(row, structure.columns["org_description"])
            )

            if park:
                result.parks.add(park)
            if department:
                result.departments.add(department)
            if organization:
                result.organizations.add(organization)

            days = parse_days(
                cell_value(row, structure.columns["days_until_due"])
            )
            if days is None or days > DUE_SOON_DAYS:
                continue

            item_title = clean_text(
                cell_value(row, structure.columns["item_title"])
            )
            status = "Overdue" if days < 0 else "Due Soon"
            key = item_title, days, status, organization

            if key in seen_pending:
                continue

            seen_pending.add(key)
            result.pending_items.append(
                PendingItem(
                    item_title=item_title,
                    days_until_due=days,
                    status=status,
                    organization=organization,
                )
            )
    finally:
        workbook.close()

    if result is not None:
        result.pending_items.sort(
            key=lambda item: (
                0 if item.status == "Overdue" else 1,
                item.days_until_due,
                item.item_title.casefold(),
            )
        )

    return result


class SSOSearchWindow(tk.Toplevel):
    STATUS_COLORS = {
        "Good": ("#C6E0B4", "#1F4E24"),
        "Due Soon": ("#FFE699", "#7F6000"),
        "Overdue": ("#F4CCCC", "#9C0006"),
    }

    def __init__(self, parent: tk.Misc, structure: WorkbookStructure) -> None:
        super().__init__(parent)
        self.structure = structure
        self.title("Search SSO - Success Factors Weekly Training Report")
        self.geometry("920x620")
        self.minsize(820, 560)
        self.transient(parent)

        self.sso_var = tk.StringVar()
        self.name_var = tk.StringVar(value="Technician: -")
        self.employee_var = tk.StringVar(value="SSO / Employee ID: -")
        self.park_var = tk.StringVar(value="Park: -")
        self.department_var = tk.StringVar(value="Department: -")
        self.organization_var = tk.StringVar(value="Organization / Venue: -")
        self.status_var = tk.StringVar(value="Enter an SSO / Employee ID to search.")

        self._build_ui()
        self.sso_entry.focus_set()
        self.bind("<Return>", lambda _event: self._start_search())

    def _build_ui(self) -> None:
        container = ttk.Frame(self, padding=16)
        container.pack(fill="both", expand=True)
        container.columnconfigure(0, weight=1)
        container.rowconfigure(3, weight=1)

        ttk.Label(
            container,
            text="Search Employee Training Status",
            font=("Segoe UI", 17, "bold"),
        ).grid(row=0, column=0, sticky="w")

        search_frame = ttk.LabelFrame(container, text="SSO / Employee ID", padding=10)
        search_frame.grid(row=1, column=0, sticky="ew", pady=(12, 10))
        search_frame.columnconfigure(0, weight=1)

        self.sso_entry = ttk.Entry(
            search_frame,
            textvariable=self.sso_var,
            font=("Segoe UI", 11),
        )
        self.sso_entry.grid(row=0, column=0, sticky="ew", padx=(0, 8))

        self.search_button = ttk.Button(
            search_frame,
            text="Search",
            command=self._start_search,
        )
        self.search_button.grid(row=0, column=1)

        details = ttk.LabelFrame(container, text="Employee Status", padding=10)
        details.grid(row=2, column=0, sticky="ew", pady=(0, 10))
        details.columnconfigure(0, weight=1)

        ttk.Label(details, textvariable=self.name_var, font=("Segoe UI", 11, "bold")).grid(
            row=0, column=0, sticky="w"
        )
        ttk.Label(details, textvariable=self.employee_var).grid(row=1, column=0, sticky="w")
        ttk.Label(details, textvariable=self.park_var).grid(row=2, column=0, sticky="w")
        ttk.Label(details, textvariable=self.department_var).grid(row=3, column=0, sticky="w")
        ttk.Label(details, textvariable=self.organization_var, wraplength=820).grid(
            row=4, column=0, sticky="w"
        )

        self.status_label = tk.Label(
            details,
            textvariable=self.status_var,
            font=("Segoe UI", 12, "bold"),
            padx=10,
            pady=6,
            anchor="w",
            relief="solid",
            borderwidth=1,
        )
        self.status_label.grid(row=5, column=0, sticky="ew", pady=(8, 0))

        results_frame = ttk.LabelFrame(
            container,
            text=f"Pending Training (Days Until Due <= {DUE_SOON_DAYS})",
            padding=10,
        )
        results_frame.grid(row=3, column=0, sticky="nsew")
        results_frame.columnconfigure(0, weight=1)
        results_frame.rowconfigure(0, weight=1)

        columns = ("item", "days", "status", "organization")
        self.tree = ttk.Treeview(
            results_frame,
            columns=columns,
            show="headings",
            selectmode="browse",
        )
        self.tree.heading("item", text="Item Title")
        self.tree.heading("days", text="Days Until Due")
        self.tree.heading("status", text="Status")
        self.tree.heading("organization", text="Organization / Venue")
        self.tree.column("item", width=390, minwidth=250, anchor="w")
        self.tree.column("days", width=110, minwidth=95, anchor="center")
        self.tree.column("status", width=100, minwidth=85, anchor="center")
        self.tree.column("organization", width=230, minwidth=170, anchor="w")
        self.tree.grid(row=0, column=0, sticky="nsew")

        scrollbar = ttk.Scrollbar(results_frame, orient="vertical", command=self.tree.yview)
        scrollbar.grid(row=0, column=1, sticky="ns")
        self.tree.configure(yscrollcommand=scrollbar.set)

        self.tree.tag_configure("Overdue", background="#F4CCCC")
        self.tree.tag_configure("Due Soon", background="#FFE699")

        self.progress = ttk.Progressbar(container, mode="indeterminate")
        self.progress.grid(row=4, column=0, sticky="ew", pady=(10, 0))

    def _start_search(self) -> None:
        employee_id = self.sso_var.get().strip()
        if not employee_id:
            messagebox.showwarning(
                "Search SSO",
                "Enter an SSO / Employee ID.",
                parent=self,
            )
            self.sso_entry.focus_set()
            return

        self._set_busy(True)
        self._clear_results()
        self.status_var.set("Searching the loaded Weekly Training report...")
        self._set_status_color(None)

        thread = threading.Thread(
            target=self._search_worker,
            args=(employee_id,),
            daemon=True,
        )
        thread.start()

    def _search_worker(self, employee_id: str) -> None:
        try:
            result = lookup_employee_status(self.structure, employee_id)
        except Exception as exc:
            self.after(0, self._search_failed, str(exc))
            return

        self.after(0, self._search_complete, employee_id, result)

    def _search_failed(self, error_text: str) -> None:
        self._set_busy(False)
        self.status_var.set("The search could not be completed.")
        self._set_status_color(None)
        messagebox.showerror("Search SSO", error_text, parent=self)

    def _search_complete(
        self,
        employee_id: str,
        result: EmployeeStatus | None,
    ) -> None:
        self._set_busy(False)

        if result is None:
            self.name_var.set("Technician: Not found")
            self.employee_var.set(f"SSO / Employee ID: {employee_id}")
            self.park_var.set("Park: -")
            self.department_var.set("Department: -")
            self.organization_var.set("Organization / Venue: -")
            self.status_var.set("No employee was found for this SSO / Employee ID.")
            self._set_status_color(None)
            return

        self.name_var.set(f"Technician: {result.employee_name or '-'}")
        self.employee_var.set(f"SSO / Employee ID: {result.employee_id}")
        self.park_var.set(f"Park: {self._join_values(result.parks)}")
        self.department_var.set(
            f"Department: {self._join_values(result.departments)}"
        )
        self.organization_var.set(
            f"Organization / Venue: {self._join_values(result.organizations)}"
        )

        overall = result.overall_status
        if overall == "Good":
            self.status_var.set(
                "Overall Status: GOOD - No overdue or due-soon training items found."
            )
        elif overall == "Due Soon":
            self.status_var.set(
                f"Overall Status: DUE SOON - {len(result.pending_items)} pending item(s)."
            )
        else:
            self.status_var.set(
                f"Overall Status: OVERDUE - {len(result.pending_items)} pending item(s)."
            )

        self._set_status_color(overall)

        for item in result.pending_items:
            self.tree.insert(
                "",
                tk.END,
                values=(
                    item.item_title,
                    item.days_until_due,
                    item.status,
                    item.organization,
                ),
                tags=(item.status,),
            )

    def _clear_results(self) -> None:
        for item_id in self.tree.get_children():
            self.tree.delete(item_id)

        self.name_var.set("Technician: -")
        self.employee_var.set("SSO / Employee ID: -")
        self.park_var.set("Park: -")
        self.department_var.set("Department: -")
        self.organization_var.set("Organization / Venue: -")

    def _set_busy(self, busy: bool) -> None:
        if busy:
            self.progress.start(10)
            self.search_button.configure(state="disabled")
            self.sso_entry.configure(state="disabled")
        else:
            self.progress.stop()
            self.search_button.configure(state="normal")
            self.sso_entry.configure(state="normal")
            self.sso_entry.focus_set()

    def _set_status_color(self, status: str | None) -> None:
        fill, foreground = self.STATUS_COLORS.get(
            status,
            ("#E7E6E6", "#404040"),
        )
        self.status_label.configure(background=fill, foreground=foreground)

    @staticmethod
    def _join_values(values: set[str]) -> str:
        return ", ".join(sorted(values)) if values else "-"
