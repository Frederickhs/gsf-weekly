from __future__ import annotations

import os
import re
import shutil
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension

from input_scanner import WorkbookStructure, cell_value, clean_text


PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_FOLDER = PROJECT_ROOT / "Output"
ARCHIVE_FOLDER = PROJECT_ROOT / "Archive"
DUE_SOON_DAYS = 15

DARK_BLUE = "17365D"
MEDIUM_BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
WHITE = "FFFFFF"
LIGHT_GRAY = "E7E6E6"
GOOD_GREEN = "C6E0B4"
DUE_YELLOW = "FFE699"
OVERDUE_RED = "F4CCCC"
BORDER_COLOR = "B7B7B7"

THIN_BORDER = Border(
    left=Side(style="thin", color=BORDER_COLOR),
    right=Side(style="thin", color=BORDER_COLOR),
    top=Side(style="thin", color=BORDER_COLOR),
    bottom=Side(style="thin", color=BORDER_COLOR),
)


@dataclass(frozen=True)
class TrainingRecord:
    employee_id: str
    display_name: str
    organization: str
    item_title: str
    days_until_due: int | float | None
    status: str


@dataclass
class OrganizationReport:
    organization: str
    employee_ids: set[str]
    records: list[TrainingRecord]
    good_employee_ids: set[str]

    @property
    def overdue_employee_ids(self) -> set[str]:
        return {
            record.employee_id
            for record in self.records
            if record.status == "Overdue"
        }

    @property
    def due_soon_employee_ids(self) -> set[str]:
        return {
            record.employee_id
            for record in self.records
            if record.status == "Due Soon"
        }


@dataclass
class GenerationResult:
    output_path: Path
    organization_count: int
    employee_count: int
    action_item_count: int


def parse_days(value: object) -> int | float | None:
    if value is None or value == "":
        return None

    try:
        number = float(value)
    except (TypeError, ValueError):
        return None

    if number.is_integer():
        return int(number)
    return number


def status_for_days(days: int | float) -> str:
    return "Overdue" if days < 0 else "Due Soon"


def technician_name(row: tuple, structure: WorkbookStructure) -> str:
    first_column = structure.columns.get("first_name")
    last_column = structure.columns.get("last_name")

    first_name = clean_text(cell_value(row, first_column)) if first_column else ""
    last_name = clean_text(cell_value(row, last_column)) if last_column else ""
    full_name = " ".join(part for part in (first_name, last_name) if part).strip()

    if full_name:
        return full_name

    return clean_text(cell_value(row, structure.columns["display_name"]))


def collect_report_data(
    structure: WorkbookStructure,
    park: str,
    department: str,
    selected_organizations: set[str],
) -> dict[str, OrganizationReport]:
    workbook = load_workbook(
        filename=structure.file_path,
        read_only=True,
        data_only=True,
    )

    employee_names: dict[tuple[str, str], str] = {}
    employee_ids_by_org: dict[str, set[str]] = defaultdict(set)
    records_by_org: dict[str, list[TrainingRecord]] = defaultdict(list)

    try:
        worksheet = workbook[structure.worksheet_name]

        for row in worksheet.iter_rows(
            min_row=structure.header_row + 1,
            values_only=True,
        ):
            row_park = clean_text(
                cell_value(row, structure.columns["park_description"])
            )
            row_department = clean_text(
                cell_value(row, structure.columns["department"])
            )
            organization = clean_text(
                cell_value(row, structure.columns["org_description"])
            )

            if (
                row_park != park
                or row_department != department
                or organization not in selected_organizations
            ):
                continue

            employee_id = clean_text(
                cell_value(row, structure.columns["employee_id"])
            )
            if not employee_id:
                continue

            display_name = technician_name(row, structure)
            employee_names[(organization, employee_id)] = display_name
            employee_ids_by_org[organization].add(employee_id)

            days = parse_days(
                cell_value(row, structure.columns["days_until_due"])
            )
            if days is None or days > DUE_SOON_DAYS:
                continue

            item_title = clean_text(
                cell_value(row, structure.columns["item_title"])
            )
            records_by_org[organization].append(
                TrainingRecord(
                    employee_id=employee_id,
                    display_name=display_name,
                    organization=organization,
                    item_title=item_title,
                    days_until_due=days,
                    status=status_for_days(days),
                )
            )
    finally:
        workbook.close()

    reports: dict[str, OrganizationReport] = {}

    for organization in sorted(selected_organizations):
        employee_ids = employee_ids_by_org.get(organization, set())
        action_employee_ids = {
            record.employee_id for record in records_by_org.get(organization, [])
        }
        good_ids = employee_ids - action_employee_ids

        good_records = [
            TrainingRecord(
                employee_id=employee_id,
                display_name=employee_names.get((organization, employee_id), ""),
                organization=organization,
                item_title="",
                days_until_due=None,
                status="Good",
            )
            for employee_id in good_ids
        ]

        all_records = records_by_org.get(organization, []) + good_records
        all_records.sort(
            key=lambda record: (
                record.display_name.casefold(),
                record.employee_id,
                0 if record.status == "Overdue" else 1 if record.status == "Due Soon" else 2,
                record.days_until_due if record.days_until_due is not None else 999999,
                record.item_title.casefold(),
            )
        )

        reports[organization] = OrganizationReport(
            organization=organization,
            employee_ids=set(employee_ids),
            records=all_records,
            good_employee_ids=set(good_ids),
        )

    return reports


def safe_sheet_name(value: str, used_names: set[str]) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "-", value).strip()
    cleaned = re.sub(r"\s+", " ", cleaned) or "Organization"
    base = cleaned[:31]
    candidate = base
    number = 2

    while candidate.casefold() in used_names:
        suffix = f" ({number})"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        number += 1

    used_names.add(candidate.casefold())
    return candidate


def style_title(ws, title: str, subtitle: str, end_column: int) -> None:
    end_letter = get_column_letter(end_column)
    ws.merge_cells(f"A1:{end_letter}1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Aptos Display", size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=DARK_BLUE)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells(f"A2:{end_letter}2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Aptos", size=10, italic=True, color="404040")
    ws["A2"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 22


def style_header_row(ws, row_number: int, end_column: int) -> None:
    for cell in ws[row_number][:end_column]:
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=MEDIUM_BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
    ws.row_dimensions[row_number].height = 22


def add_status_fill(cell, status: str) -> None:
    color = {
        "Good": GOOD_GREEN,
        "Due Soon": DUE_YELLOW,
        "Overdue": OVERDUE_RED,
    }.get(status)
    if color:
        cell.fill = PatternFill("solid", fgColor=color)
    cell.font = Font(name="Aptos", size=10, bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def create_summary_sheet(
    workbook: Workbook,
    reports: dict[str, OrganizationReport],
    park: str,
    department: str,
    generated_at: datetime,
) -> None:
    ws = workbook.active
    ws.title = "Summary"

    style_title(
        ws,
        "Learning Weekly Training Report",
        f"Park: {park} | Department: {department} | Generated: {generated_at:%m/%d/%Y %I:%M %p}",
        6,
    )

    headers = [
        "Organization / Venue",
        "Employees",
        "Good",
        "Due Soon",
        "Overdue",
        "Needs Attention",
    ]
    for column, header in enumerate(headers, start=1):
        ws.cell(row=4, column=column, value=header)
    style_header_row(ws, 4, len(headers))

    total_employees: set[str] = set()
    total_good = 0
    total_due_soon = 0
    total_overdue = 0
    total_action_items = 0

    for row_number, organization in enumerate(sorted(reports), start=5):
        report = reports[organization]
        due_soon = report.due_soon_employee_ids
        overdue = report.overdue_employee_ids
        needs_attention = due_soon | overdue
        action_items = sum(
            1 for record in report.records if record.status in {"Due Soon", "Overdue"}
        )

        values = [
            organization,
            len(report.employee_ids),
            len(report.good_employee_ids),
            len(due_soon),
            len(overdue),
            len(needs_attention),
        ]

        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=row_number, column=column, value=value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(
                horizontal="left" if column == 1 else "center",
                vertical="center",
            )

        ws.cell(row=row_number, column=3).fill = PatternFill("solid", fgColor=GOOD_GREEN)
        ws.cell(row=row_number, column=4).fill = PatternFill("solid", fgColor=DUE_YELLOW)
        ws.cell(row=row_number, column=5).fill = PatternFill("solid", fgColor=OVERDUE_RED)

        total_employees.update(report.employee_ids)
        total_good += len(report.good_employee_ids)
        total_due_soon += len(due_soon)
        total_overdue += len(overdue)
        total_action_items += action_items

    total_row = 5 + len(reports)
    total_needs_attention = len(
        {
            record.employee_id
            for report in reports.values()
            for record in report.records
            if record.status in {"Due Soon", "Overdue"}
        }
    )
    totals = [
        "TOTAL",
        len(total_employees),
        total_good,
        total_due_soon,
        total_overdue,
        total_needs_attention,
    ]

    for column, value in enumerate(totals, start=1):
        cell = ws.cell(row=total_row, column=column, value=value)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=DARK_BLUE)
        cell.border = THIN_BORDER
        cell.alignment = Alignment(
            horizontal="left" if column == 1 else "center",
            vertical="center",
        )

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:F{max(4, total_row - 1)}"
    widths = {"A": 52, "B": 12, "C": 12, "D": 12, "E": 12, "F": 16}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.sheet_view.showGridLines = False


def create_organization_sheet(
    workbook: Workbook,
    report: OrganizationReport,
    park: str,
    department: str,
    generated_at: datetime,
    used_names: set[str],
) -> None:
    title = safe_sheet_name(report.organization, used_names)
    ws = workbook.create_sheet(title)

    style_title(
        ws,
        report.organization,
        f"Park: {park} | Department: {department} | Days Until Due <= {DUE_SOON_DAYS} | Generated: {generated_at:%m/%d/%Y %I:%M %p}",
        5,
    )

    headers = ["Technician", "SSO", "Item Title", "Days Until Due", "Status"]
    for column, header in enumerate(headers, start=1):
        ws.cell(row=4, column=column, value=header)
    style_header_row(ws, 4, len(headers))

    for row_number, record in enumerate(report.records, start=5):
        values = [
            record.display_name,
            record.employee_id,
            record.item_title,
            record.days_until_due,
            record.status,
        ]

        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=row_number, column=column, value=value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(
                horizontal="center" if column in {2, 4, 5} else "left",
                vertical="top",
                wrap_text=column == 3,
            )

        ws.cell(row=row_number, column=2).number_format = "@"
        add_status_fill(ws.cell(row=row_number, column=5), record.status)

    final_row = max(5, 4 + len(report.records))
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:E{final_row}"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 70
    ws.column_dimensions["D"].width = 17
    ws.column_dimensions["E"].width = 14
    ws.sheet_view.showGridLines = False


def archive_existing_output(
    output_folder: Path,
    archive_folder: Path = ARCHIVE_FOLDER,
) -> Path | None:
    output_folder.mkdir(parents=True, exist_ok=True)
    existing_items = [item for item in output_folder.iterdir()]

    if not existing_items:
        return None

    timestamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    destination = archive_folder / f"Output_{timestamp}"
    counter = 2

    while destination.exists():
        destination = archive_folder / f"Output_{timestamp}_{counter}"
        counter += 1

    destination.mkdir(parents=True, exist_ok=False)

    for item in existing_items:
        shutil.move(str(item), str(destination / item.name))

    return destination


def generate_report(
    structure: WorkbookStructure,
    park: str,
    department: str,
    selected_organizations: set[str],
    output_folder: Path | None = None,
) -> GenerationResult:
    if not selected_organizations:
        raise ValueError("At least one Organization / Venue must be selected.")

    reports = collect_report_data(
        structure=structure,
        park=park,
        department=department,
        selected_organizations=selected_organizations,
    )

    generated_at = datetime.now()
    target_folder = output_folder or OUTPUT_FOLDER
    archive_existing_output(target_folder)
    target_folder.mkdir(parents=True, exist_ok=True)

    filename = f"Learning Weekly Training Report {generated_at:%Y-%m-%d_%H%M%S}.xlsx"
    output_path = target_folder / filename

    workbook = Workbook()
    create_summary_sheet(
        workbook=workbook,
        reports=reports,
        park=park,
        department=department,
        generated_at=generated_at,
    )

    used_names = {"summary"}
    for organization in sorted(reports):
        create_organization_sheet(
            workbook=workbook,
            report=reports[organization],
            park=park,
            department=department,
            generated_at=generated_at,
            used_names=used_names,
        )

    workbook.save(output_path)

    unique_employees = {
        employee_id
        for report in reports.values()
        for employee_id in report.employee_ids
    }
    action_item_count = sum(
        1
        for report in reports.values()
        for record in report.records
        if record.status in {"Due Soon", "Overdue"}
    )

    return GenerationResult(
        output_path=output_path,
        organization_count=len(reports),
        employee_count=len(unique_employees),
        action_item_count=action_item_count,
    )


def open_output(path: Path) -> None:
    if hasattr(os, "startfile"):
        os.startfile(path)  # type: ignore[attr-defined]
