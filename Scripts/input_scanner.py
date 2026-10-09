from __future__ import annotations

import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = (
    Path(sys.executable).resolve().parent
    if getattr(sys, "frozen", False)
    else Path(__file__).resolve().parent.parent
)
INPUT_FOLDER = PROJECT_ROOT / "Input"
SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}

REQUIRED_COLUMNS = {
    "employee_id": {
        "employee id",
        "employee id/sso",
        "employee number",
        "sso",
        "labor id",
        "laborid",
    },
    "first_name": {"first name", "first name ", "f.name", "firstname"},
    "last_name": {"last name", "last name ", "l.name", "lastname"},
    "display_name": {
        "display name",
        "employee name",
        "technician",
        "technician name",
        "team member",
        "team member name",
    },
    "item_title": {
        "item title",
        "training title",
        "course title",
        "item name",
    },
    "days_until_due": {
        "days until due",
        "days to due",
        "days remaining",
    },
    "park_description": {
        "park description",
        "park",
        "property",
    },
    "org_description": {
        "org description",
        "organization description",
        "organization",
        "org",
        "venue",
    },
    "department": {
        "sub division description",
        "department",
        "department description",
        "subdivision description",
    },
}


@dataclass
class WorkbookStructure:
    file_path: Path
    worksheet_name: str
    header_row: int
    columns: dict[str, int]
    actual_headers: dict[str, str]


@dataclass
class FilterResults:
    parks: dict[str, set[str]]
    departments: dict[str, set[str]]
    organizations: dict[tuple[str, str], set[str]]
    employees_by_organization: dict[tuple[str, str, str], set[str]]
    total_data_rows: int


def normalize_header(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().lower().split())


def find_matching_column(normalized_header: str) -> str | None:
    for field_name, aliases in REQUIRED_COLUMNS.items():
        if normalized_header in aliases:
            return field_name
    return None


def find_input_workbooks(folder: Path) -> list[Path]:
    if not folder.exists():
        return []

    candidates = [
        path
        for path in folder.iterdir()
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_EXTENSIONS
        and not path.name.startswith("~$")
    ]

    return sorted(
        candidates,
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )


def inspect_worksheet_headers(
    worksheet,
    maximum_rows: int = 25,
) -> tuple[int, dict[str, int], dict[str, str]] | None:
    best_result = None
    best_score = 0

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            max_row=maximum_rows,
            values_only=True,
        ),
        start=1,
    ):
        mapped_columns: dict[str, int] = {}
        actual_headers: dict[str, str] = {}

        for column_number, value in enumerate(row, start=1):
            normalized = normalize_header(value)
            field_name = find_matching_column(normalized)

            if field_name and field_name not in mapped_columns:
                mapped_columns[field_name] = column_number
                actual_headers[field_name] = str(value).strip()

        score = len(mapped_columns)
        if score > best_score:
            best_score = score
            best_result = row_number, mapped_columns, actual_headers

    if best_result is None:
        return None

    header_row, columns, actual_headers = best_result
    minimum_required = {
        "employee_id",
        "display_name",
        "item_title",
        "days_until_due",
        "park_description",
        "org_description",
        "department",
    }

    if not minimum_required.issubset(columns):
        return None

    return header_row, columns, actual_headers


def inspect_workbook(file_path: Path) -> WorkbookStructure:
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    try:
        for worksheet in workbook.worksheets:
            result = inspect_worksheet_headers(worksheet)
            if result is None:
                continue

            header_row, columns, actual_headers = result
            return WorkbookStructure(
                file_path=file_path,
                worksheet_name=worksheet.title,
                header_row=header_row,
                columns=columns,
                actual_headers=actual_headers,
            )
    finally:
        workbook.close()

    raise ValueError(
        "No worksheet containing all required Weekly Training columns was found."
    )


def cell_value(row: tuple, column_number: int) -> object:
    index = column_number - 1
    if index < 0 or index >= len(row):
        return None
    return row[index]


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def scan_filters(structure: WorkbookStructure) -> FilterResults:
    workbook = load_workbook(
        filename=structure.file_path,
        read_only=True,
        data_only=True,
    )

    parks: dict[str, set[str]] = {}
    departments: dict[str, set[str]] = {}
    organizations: dict[tuple[str, str], set[str]] = {}
    employees_by_organization: dict[tuple[str, str, str], set[str]] = {}
    total_data_rows = 0

    try:
        worksheet = workbook[structure.worksheet_name]

        for row in worksheet.iter_rows(
            min_row=structure.header_row + 1,
            values_only=True,
        ):
            employee_id = clean_text(
                cell_value(row, structure.columns["employee_id"])
            )
            park = clean_text(
                cell_value(row, structure.columns["park_description"])
            )
            department = clean_text(
                cell_value(row, structure.columns["department"])
            )
            organization = clean_text(
                cell_value(row, structure.columns["org_description"])
            )

            if not employee_id:
                continue

            total_data_rows += 1

            if park:
                parks.setdefault(park, set()).add(employee_id)

            if park and department:
                departments.setdefault(park, set()).add(department)

            if park and department and organization:
                organizations.setdefault((park, department), set()).add(
                    organization
                )
                employees_by_organization.setdefault(
                    (park, department, organization), set()
                ).add(employee_id)
    finally:
        workbook.close()

    return FilterResults(
        parks=parks,
        departments=departments,
        organizations=organizations,
        employees_by_organization=employees_by_organization,
        total_data_rows=total_data_rows,
    )


def print_heading(text: str) -> None:
    print()
    print("=" * 72)
    print(text)
    print("=" * 72)


def print_detected_columns(structure: WorkbookStructure) -> None:
    print_heading("DETECTED COLUMNS")

    for field_name, column_number in structure.columns.items():
        actual_header = structure.actual_headers[field_name]
        print(
            f"{field_name:<22} "
            f"Column {column_number:<3} "
            f"{actual_header}"
        )


def print_filter_summary(results: FilterResults) -> None:
    print_heading("PARKS")

    for park in sorted(results.parks):
        employee_count = len(results.parks[park])
        department_count = len(results.departments.get(park, set()))
        print(
            f"{park} | "
            f"{employee_count} unique employees | "
            f"{department_count} departments"
        )

    print_heading("PARK / DEPARTMENT / ORGANIZATION")

    for park in sorted(results.departments):
        for department in sorted(results.departments[park]):
            key = park, department
            organization_values = sorted(
                results.organizations.get(key, set())
            )

            print()
            print(f"Park: {park}")
            print(f"Department: {department}")
            print(f"Organizations: {len(organization_values)}")

            for organization in organization_values:
                employee_count = len(
                    results.employees_by_organization.get(
                        (park, department, organization), set()
                    )
                )
                print(
                    f"  - {organization} "
                    f"({employee_count} unique employees)"
                )


def main() -> int:
    print_heading("SUCCESS FACTORS WEEKLY TRAINING REPORT - INPUT SCANNER")

    workbooks = find_input_workbooks(INPUT_FOLDER)
    if not workbooks:
        print("No supported Weekly Training workbook was found in:")
        print(INPUT_FOLDER)
        print()
        print("Place one .xlsx or .xlsm file in the Input folder.")
        return 1

    input_file = workbooks[0]
    print(f"Input file: {input_file.name}")

    if len(workbooks) > 1:
        print(
            "Multiple workbooks were found. "
            "The most recently modified workbook was selected."
        )

    try:
        structure = inspect_workbook(input_file)
    except Exception as exc:
        print()
        print(f"ERROR: {exc}")
        return 1

    print(f"Worksheet: {structure.worksheet_name}")
    print(f"Header row: {structure.header_row}")
    print_detected_columns(structure)

    try:
        results = scan_filters(structure)
    except Exception as exc:
        print()
        print(f"ERROR WHILE READING DATA: {exc}")
        return 1

    print_filter_summary(results)

    print_heading("SCAN COMPLETE")
    print(f"Rows read: {results.total_data_rows:,}")
    print(f"Parks detected: {len(results.parks):,}")
    print(
        "The Weekly Training workbook is ready for "
        "Park, Department, and Organization filtering."
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
